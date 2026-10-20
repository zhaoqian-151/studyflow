import hashlib
import io
import math
import re
import zipfile
from collections import Counter

from fastapi import HTTPException
from sqlalchemy import delete, select, text
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Chunk, Document
from app.services.ai import embed


def split_text(text: str, size=600, overlap=80) -> list[str]:
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("Chunk overlap must be smaller than size")
    text = text.replace("\r\n", "\n").strip()
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + size // 2, end), text.rfind("。", start + size // 2, end))
            if boundary >= 0:
                end = boundary + 1
        value = text[start:end].strip()
        if value:
            chunks.append(value)
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def tokenize(text: str):
    text = text.lower()
    terms = re.findall(r"[a-z0-9_]+", text)
    for run in re.findall(r"[\u4e00-\u9fff]+", text):
        terms.extend(run)
        terms.extend(run[i : i + 2] for i in range(len(run) - 1))
    return terms


def bm25_scores(query: str, texts: list[str]):
    tokens = [Counter(tokenize(text)) for text in texts]
    lengths = [sum(item.values()) for item in tokens]
    avg = sum(lengths) / max(len(texts), 1) or 1
    scores = [0.0] * len(texts)
    for term in set(tokenize(query)):
        df = sum(term in item for item in tokens)
        inverse = math.log(1 + (len(texts) - df + 0.5) / (df + 0.5))
        for i, item in enumerate(tokens):
            freq = item[term]
            scores[i] += inverse * freq * 2.5 / (freq + 1.5 * (0.25 + 0.75 * lengths[i] / avg))
    return scores


def cosine(a, b):
    if len(a) != len(b):
        return 0
    denominator = math.sqrt(sum(v * v for v in a) * sum(v * v for v in b))
    return sum(x * y for x, y in zip(a, b)) / denominator if denominator else 0


def extract_file(filename: str, content: bytes):
    suffix = filename.lower().rsplit(".", 1)[-1]
    try:
        if suffix in {"txt", "md"}:
            return content.decode("utf-8-sig")
        if suffix == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if len(reader.pages) > 150:
                raise HTTPException(413, "PDF 最多 150 页")
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix == "docx":
            from docx import Document as WordDocument

            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise HTTPException(413, "DOCX 解压体积过大")
            doc = WordDocument(io.BytesIO(content))
            return "\n".join(
                [p.text for p in doc.paragraphs]
                + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
            )
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, "文档解析失败，请使用 UTF-8 文本或未加密的 PDF / DOCX") from error
    raise HTTPException(415, "仅支持 .md / .txt / .pdf / .docx")


async def ingest(db: Session, user_id: str, title: str, content: str):
    content = content.strip()
    if not content:
        raise HTTPException(422, "未提取到文本；扫描版 PDF 请先进行 OCR")
    if len(content) > 200000:
        raise HTTPException(413, "文档文本最多 20 万字")
    digest = hashlib.sha256(content.encode()).hexdigest()
    existing = db.scalar(select(Document).where(Document.user_id == user_id, Document.digest == digest))
    if existing:
        return {"id": existing.id, "duplicate": True}
    pieces = split_text(content)
    vectors = await embed(pieces) if settings.embedding_api_key else None
    doc = Document(user_id=user_id, title=title, content=content, digest=digest)
    db.add(doc)
    db.flush()
    for i, piece in enumerate(pieces):
        db.add(
            Chunk(
                user_id=user_id,
                document_id=doc.id,
                position=i,
                content=piece,
                embedding=vectors[i] if vectors else None,
                embedding_model=settings.embedding_model if vectors else None,
            )
        )
    db.commit()
    return {"id": doc.id, "chunks": len(pieces), "duplicate": False, "index_mode": "hybrid" if vectors else "bm25"}


async def reindex_document(db: Session, document: Document):
    if not settings.embedding_api_key:
        raise HTTPException(409, "请先配置向量模型，再重建索引")
    pieces = split_text(document.content)
    # Provider failures occur before any replacement; the original index remains usable.
    vectors = await embed(pieces)
    locked = db.scalar(
        select(Document).where(Document.id == document.id, Document.user_id == document.user_id).with_for_update()
    )
    if not locked:
        raise HTTPException(404, "文档已被删除")
    db.execute(delete(Chunk).where(Chunk.document_id == document.id, Chunk.user_id == document.user_id))
    for position, piece in enumerate(pieces):
        db.add(
            Chunk(
                user_id=document.user_id,
                document_id=document.id,
                position=position,
                content=piece,
                embedding=vectors[position],
                embedding_model=settings.embedding_model,
            )
        )
    db.commit()
    return {"id": document.id, "chunks": len(pieces), "index_mode": "hybrid", "index_status": "ready"}


def vector_statement(user_id: str, vector: list, candidates=20, document_id=None):
    distance = Chunk.embedding.cosine_distance(vector)
    statement = (
        select(Chunk.id, distance.label("distance"))
        .where(
            Chunk.user_id == user_id, Chunk.embedding_model == settings.embedding_model, Chunk.embedding.is_not(None)
        )
        .order_by(distance)
        .limit(candidates)
    )
    return statement.where(Chunk.document_id == document_id) if document_id else statement


def semantic_candidates(db: Session, user_id: str, vector: list, candidates=20, document_id=None):
    if db.bind.dialect.name == "postgresql":
        # PostgreSQL returns only IDs and distances. Stored vectors never leave the database.
        # Iterative scans compensate for filtering by user/model after HNSW traversal.
        with db.begin_nested():
            db.execute(text("SET LOCAL hnsw.iterative_scan = 'strict_order'"))
            rows = db.execute(vector_statement(user_id, vector, candidates, document_id)).all()
        return [(row.id, 1 - row.distance) for row in rows]
    # Portable test/development fallback; production uses pgvector above.
    statement = select(Chunk.id, Chunk.embedding).where(
        Chunk.user_id == user_id, Chunk.embedding_model == settings.embedding_model, Chunk.embedding.is_not(None)
    )
    if document_id:
        statement = statement.where(Chunk.document_id == document_id)
    rows = db.execute(statement).all()
    return sorted([(row.id, cosine(vector, row.embedding)) for row in rows], key=lambda item: item[1], reverse=True)[
        :candidates
    ]


async def search(db: Session, user_id: str, query: str, limit=4, document_id=None):
    # BM25 currently reads this user's text corpus, but never selects embedding columns.
    statement = (
        select(Chunk.id, Chunk.document_id, Chunk.position, Chunk.content, Document.title)
        .join(Document, Document.id == Chunk.document_id)
        .where(Chunk.user_id == user_id, Document.user_id == user_id)
        .order_by(Chunk.id)
    )
    if document_id:
        statement = statement.where(Chunk.document_id == document_id)
    rows = db.execute(statement).all()
    if not rows:
        return [], "bm25"
    by_id = {row.id: row for row in rows}
    lexical = bm25_scores(query, [row.content for row in rows])
    candidates = 20
    lexical_order = sorted(range(len(rows)), key=lambda i: lexical[i], reverse=True)[:candidates]
    ranking = {rows[i].id: 1 / (60 + rank) for rank, i in enumerate(lexical_order, 1) if lexical[i] > 0}
    mode, warning = "bm25", None
    if settings.embedding_api_key:
        try:
            vector = (await embed([query]))[0]
            semantic = semantic_candidates(db, user_id, vector, candidates, document_id)
            for rank, (identifier, score) in enumerate(semantic, 1):
                if score > 0.25 and identifier in by_id:
                    ranking[identifier] = ranking.get(identifier, 0) + 1 / (60 + rank)
            mode = "hybrid" if semantic else "bm25"
        except Exception:
            warning = "向量服务不可用，本次使用 BM25 检索"
    selected = sorted(ranking, key=ranking.get, reverse=True)[:limit]
    return [
        {
            "index": rank,
            "document_id": by_id[identifier].document_id,
            "title": by_id[identifier].title,
            "chunk": by_id[identifier].position,
            "content": by_id[identifier].content,
            "score": round(ranking[identifier], 5),
            "warning": warning,
        }
        for rank, identifier in enumerate(selected, 1)
    ], mode
